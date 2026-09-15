import os
import re
from fastapi import UploadFile, HTTPException
from pypdf import PdfReader
import docx

ALLOWED_EXTENSIONS = {".pdf", ".docx"}
MAX_FILE_SIZE_MB = 10

def validate_resume_file(file: UploadFile):
    filename = file.filename
    ext = os.path.splitext(filename)[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400, 
            detail=f"Invalid file extension '{ext}'. Only PDF and DOCX files are allowed."
        )

def extract_text_from_pdf(file_path_or_bytes) -> str:
    try:
        reader = PdfReader(file_path_or_bytes)
        text_parts = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text)
        return "\n".join(text_parts)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to extract text from PDF: {str(e)}")

def extract_text_from_docx(file_path_or_bytes) -> str:
    try:
        doc = docx.Document(file_path_or_bytes)
        text_parts = [para.text for para in doc.paragraphs if para.text]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        text_parts.append(cell.text)
        return "\n".join(text_parts)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to extract text from DOCX: {str(e)}")

def extract_resume_text(file: UploadFile) -> str:
    validate_resume_file(file)
    ext = os.path.splitext(file.filename)[1].lower()
    file_bytes = file.file
    
    if ext == ".pdf":
        raw_text = extract_text_from_pdf(file_bytes)
    elif ext == ".docx":
        raw_text = extract_text_from_docx(file_bytes)
    else:
        raise HTTPException(status_code=400, detail="Unsupported file format.")

    if not raw_text or len(raw_text.strip()) < 50:
        raise HTTPException(
            status_code=400,
            detail="The uploaded document contains very little or no readable text. Please ensure it is not a scanned image PDF."
        )
    return clean_text(raw_text)

def clean_text(text: str) -> str:
    if not text:
        return ""
    # Normalize whitespace
    text = re.sub(r'\r\n|\r', '\n', text)
    text = re.sub(r'[ \t]+', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()
