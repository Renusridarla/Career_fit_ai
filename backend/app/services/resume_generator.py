import os
import io
import logging
from typing import Tuple, Dict
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from app.schemas.resume import ParsedResumeData

logger = logging.getLogger(__name__)

def generate_pdf_resume(resume: ParsedResumeData, output_path: str) -> str:
    """
    Generates a clean, ATS-optimized PDF resume containing total master resume profile sections.
    """
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'HeaderName',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=22,
        textColor=colors.HexColor('#1E293B'),
        alignment=1
    )

    contact_style = ParagraphStyle(
        'HeaderContact',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#475569'),
        alignment=1
    )

    section_heading = ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=14,
        textColor=colors.HexColor('#0F172A'),
        spaceBefore=10,
        spaceAfter=4
    )

    body_style = ParagraphStyle(
        'BodyText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor('#334155')
    )

    story = []

    # 1. Header Name & Contact
    p = resume.personal
    name = p.name if p.name else "Candidate Resume"
    story.append(Paragraph(name, title_style))

    contact_parts = []
    if p.email: contact_parts.append(p.email)
    if p.phone: contact_parts.append(p.phone)
    if p.location: contact_parts.append(p.location)
    if p.linkedin: contact_parts.append(p.linkedin)
    if p.github: contact_parts.append(p.github)
    if p.portfolio: contact_parts.append(p.portfolio)

    if contact_parts:
        story.append(Spacer(1, 4))
        story.append(Paragraph(" | ".join(contact_parts), contact_style))

    story.append(Spacer(1, 8))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#CBD5E1'), spaceBefore=2, spaceAfter=8))

    # 2. Professional Summary
    if resume.summary:
        story.append(Paragraph("PROFESSIONAL SUMMARY", section_heading))
        story.append(Paragraph(resume.summary, body_style))
        story.append(Spacer(1, 6))

    # 3. Technical Skills
    s = resume.skills
    skill_rows = []
    if s.programming_languages: skill_rows.append(f"<b>Languages:</b> {', '.join(s.programming_languages)}")
    if s.frontend: skill_rows.append(f"<b>Frontend:</b> {', '.join(s.frontend)}")
    if s.backend: skill_rows.append(f"<b>Backend:</b> {', '.join(s.backend)}")
    if s.databases: skill_rows.append(f"<b>Databases:</b> {', '.join(s.databases)}")
    if s.cloud: skill_rows.append(f"<b>Cloud/DevOps:</b> {', '.join(s.cloud)}")
    if s.tools: skill_rows.append(f"<b>Tools:</b> {', '.join(s.tools)}")
    if s.frameworks: skill_rows.append(f"<b>Frameworks:</b> {', '.join(s.frameworks)}")
    if s.other: skill_rows.append(f"<b>Other:</b> {', '.join(s.other)}")

    if skill_rows:
        story.append(Paragraph("TECHNICAL SKILLS", section_heading))
        for sr in skill_rows:
            story.append(Paragraph(sr, body_style))
        story.append(Spacer(1, 6))

    # 4. Work Experience
    if resume.experience:
        story.append(Paragraph("WORK EXPERIENCE", section_heading))
        for exp in resume.experience:
            role_line = f"<b>{exp.role}</b> — {exp.company}"
            dates_line = f"<i>{exp.dates}</i>" if exp.dates else ""
            
            table_data = [[Paragraph(role_line, body_style), Paragraph(dates_line, ParagraphStyle('RightText', parent=body_style, alignment=2))]]
            t = Table(table_data, colWidths=[380, 160])
            t.setStyle(TableStyle([
                ('VALIGN', (0,0), (-1,-1), 'TOP'),
                ('LEFTPADDING', (0,0), (-1,-1), 0),
                ('RIGHTPADDING', (0,0), (-1,-1), 0),
            ]))
            story.append(t)

            for resp in exp.responsibilities:
                story.append(Paragraph(f"• {resp}", body_style))
            for ach in exp.achievements:
                story.append(Paragraph(f"• <b>Achievement:</b> {ach}", body_style))
            story.append(Spacer(1, 4))
        story.append(Spacer(1, 4))

    # 5. Projects
    if resume.projects:
        story.append(Paragraph("KEY PROJECTS", section_heading))
        for proj in resume.projects:
            proj_title = f"<b>{proj.name}</b>"
            tech_line = f"<i>({', '.join(proj.technologies)})</i>" if proj.technologies else ""
            story.append(Paragraph(f"{proj_title} {tech_line}", body_style))
            if proj.description:
                story.append(Paragraph(f"• {proj.description}", body_style))
            for ach in proj.achievements:
                story.append(Paragraph(f"• {ach}", body_style))
            story.append(Spacer(1, 4))

    # 6. Education
    if resume.education:
        story.append(Paragraph("EDUCATION", section_heading))
        for edu in resume.education:
            deg = f"<b>{edu.degree} in {edu.field}</b>" if edu.field else f"<b>{edu.degree}</b>"
            inst = f"{edu.institution} ({edu.graduation_year})" if edu.graduation_year else edu.institution
            cgpa = f" | CGPA: {edu.cgpa}" if edu.cgpa else ""
            story.append(Paragraph(f"{deg} — {inst}{cgpa}", body_style))

    # 7. Certifications
    if resume.certifications:
        story.append(Spacer(1, 4))
        story.append(Paragraph("CERTIFICATIONS", section_heading))
        for cert in resume.certifications:
            story.append(Paragraph(f"• {cert}", body_style))

    # 8. Achievements
    if resume.achievements:
        story.append(Spacer(1, 4))
        story.append(Paragraph("HONORS & ACHIEVEMENTS", section_heading))
        for ach in resume.achievements:
            story.append(Paragraph(f"• {ach}", body_style))

    doc.build(story)
    return output_path

def generate_docx_resume(resume: ParsedResumeData, output_path: str) -> str:
    """
    Generates a clean ATS DOCX resume using python-docx containing total master resume sections.
    """
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc = docx.Document()

    # Margins
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # Name Header
    p = resume.personal
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name_p.add_run(p.name if p.name else "Candidate Resume")
    run.bold = True
    run.font.size = Pt(18)
    run.font.name = 'Calibri'
    run.font.color.rgb = RGBColor(30, 41, 59)

    # Contact line
    contact_parts = [part for part in [p.email, p.phone, p.location, p.linkedin, p.github, p.portfolio] if part]
    if contact_parts:
        c_p = doc.add_paragraph()
        c_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c_run = c_p.add_run(" | ".join(contact_parts))
        c_run.font.size = Pt(9.5)
        c_run.font.name = 'Calibri'
        c_run.font.color.rgb = RGBColor(71, 85, 105)

    def add_section_header(title: str):
        h_p = doc.add_paragraph()
        h_p.paragraph_format.space_before = Pt(8)
        h_p.paragraph_format.space_after = Pt(2)
        h_run = h_p.add_run(title)
        h_run.bold = True
        h_run.font.size = Pt(12)
        h_run.font.name = 'Calibri'
        h_run.font.color.rgb = RGBColor(15, 23, 42)

    # Summary
    if resume.summary:
        add_section_header("PROFESSIONAL SUMMARY")
        s_p = doc.add_paragraph(resume.summary)
        s_p.paragraph_format.space_after = Pt(4)

    # Skills
    s = resume.skills
    skill_rows = []
    if s.programming_languages: skill_rows.append(("Languages", ", ".join(s.programming_languages)))
    if s.frontend: skill_rows.append(("Frontend", ", ".join(s.frontend)))
    if s.backend: skill_rows.append(("Backend", ", ".join(s.backend)))
    if s.databases: skill_rows.append(("Databases", ", ".join(s.databases)))
    if s.cloud: skill_rows.append(("Cloud/DevOps", ", ".join(s.cloud)))
    if s.tools: skill_rows.append(("Tools", ", ".join(s.tools)))
    if s.frameworks: skill_rows.append(("Frameworks", ", ".join(s.frameworks)))
    if s.other: skill_rows.append(("Other", ", ".join(s.other)))

    if skill_rows:
        add_section_header("TECHNICAL SKILLS")
        for label, val in skill_rows:
            sk_p = doc.add_paragraph()
            r1 = sk_p.add_run(f"{label}: ")
            r1.bold = True
            sk_p.add_run(val)

    # Experience
    if resume.experience:
        add_section_header("WORK EXPERIENCE")
        for exp in resume.experience:
            e_p = doc.add_paragraph()
            r_role = e_p.add_run(f"{exp.role} — {exp.company}")
            r_role.bold = True
            if exp.dates:
                e_p.add_run(f" ({exp.dates})").italic = True
            for resp in exp.responsibilities:
                doc.add_paragraph(f"• {resp}")

    # Projects
    if resume.projects:
        add_section_header("KEY PROJECTS")
        for proj in resume.projects:
            pr_p = doc.add_paragraph()
            r_pr = pr_p.add_run(proj.name)
            r_pr.bold = True
            if proj.technologies:
                pr_p.add_run(f" ({', '.join(proj.technologies)})").italic = True
            if proj.description:
                doc.add_paragraph(f"• {proj.description}")

    # Education
    if resume.education:
        add_section_header("EDUCATION")
        for edu in resume.education:
            ed_p = doc.add_paragraph()
            ed_run = ed_p.add_run(f"{edu.degree} in {edu.field}" if edu.field else edu.degree)
            ed_run.bold = True
            ed_p.add_run(f" — {edu.institution} ({edu.graduation_year})")

    # Certifications
    if resume.certifications:
        add_section_header("CERTIFICATIONS")
        for cert in resume.certifications:
            doc.add_paragraph(f"• {cert}")

    # Achievements
    if resume.achievements:
        add_section_header("HONORS & ACHIEVEMENTS")
        for ach in resume.achievements:
            doc.add_paragraph(f"• {ach}")

    doc.save(output_path)
    return output_path
